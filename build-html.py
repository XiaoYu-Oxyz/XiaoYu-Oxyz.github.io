#!/usr/bin/env python3
"""Build the deployable HTML from data/*.xlsx + data/设置.json.

Usage:
    python build-html.py

Rules:
  - class name = xlsx filename stem (data/Z1.xlsx -> Z1)
  - only the first worksheet of each class file is read
  - worksheet name is ignored
  - base headers: 节次 / 实验名称 / 实验地点
  - time: 日期, or 周次+星期; either group is enough, both allowed
  - per row: valid 周次+星期 wins, otherwise fall back to 日期
  - shared week1 comes from data/设置.json

The default data is written into the inert <template id="defaultData"> JSON,
so changing data does NOT change the CSP hashes. Code changes still require
verify-csp.py --apply.
"""
import json
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / "data"
SETTINGS = DATA_DIR / "设置.json"
SOURCE = BASE / "实验提醒.html"
DEPLOY = BASE / "index.html"

TEMPLATE_RE = re.compile(r'(<template id="defaultData">)([\s\S]*?)(</template>)')
REQUIRED = ["节次", "实验名称", "实验地点"]
WEEKDAY_MAP = {
    "星期一": 1, "周一": 1, "礼拜一": 1, "一": 1, "1": 1,
    "星期二": 2, "周二": 2, "礼拜二": 2, "二": 2, "2": 2,
    "星期三": 3, "周三": 3, "礼拜三": 3, "三": 3, "3": 3,
    "星期四": 4, "周四": 4, "礼拜四": 4, "四": 4, "4": 4,
    "星期五": 5, "周五": 5, "礼拜五": 5, "五": 5, "5": 5,
    "星期六": 6, "周六": 6, "礼拜六": 6, "六": 6, "6": 6,
    "星期日": 7, "星期天": 7, "周日": 7, "周天": 7,
    "礼拜日": 7, "礼拜天": 7, "日": 7, "天": 7, "7": 7, "0": 7,
}

MAX_EXPERIMENTS = 5000
MAX_PERIOD_LEN = 100
MAX_NAME_LEN = 200
MAX_PLACE_LEN = 200

SETTINGS_WEEK1 = None


def fail(message):
    raise SystemExit("BUILD FAILED: " + message)


def natural_key(name):
    key = []
    for part in re.split(r"(\d+)", name):
        if not part:
            continue
        key.append((0, int(part)) if part.isdigit() else (1, part.casefold()))
    return tuple(key)


def norm_header(value):
    return str(value).replace("\u3000", " ").strip()


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def cell_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except Exception:
            pass
    text = cell_text(value)
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def parse_week(value):
    if isinstance(value, (int, float)) and float(value).is_integer():
        number = int(value)
        return number if 1 <= number <= 1000 else None
    text = cell_text(value).replace(" ", "")
    match = re.fullmatch(r"(?:第)?(\d+)周?", text)
    if not match:
        return None
    number = int(match.group(1))
    return number if 1 <= number <= 1000 else None


def parse_weekday(value):
    if isinstance(value, (int, float)) and float(value).is_integer():
        number = int(value)
        if number == 0:
            return 7
        return number if 1 <= number <= 7 else None
    text = cell_text(value).replace(" ", "")
    return WEEKDAY_MAP.get(text)


def clean_field(value, label, max_len, required=False, default=""):
    text = cell_text(value)
    if re.search(r"[\u0000-\u0008\u000B\u000C\u000E-\u001F]", text):
        fail("%s 包含控制字符" % label)
    if required and not text:
        fail("%s 不能为空" % label)
    if len(text) > max_len:
        fail("%s 超过 %d 个字符" % (label, max_len))
    return text or default


def read_settings():
    if not SETTINGS.exists():
        fail("找不到 %s" % SETTINGS)
    try:
        data = json.loads(SETTINGS.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        fail("%s 不是有效 JSON：%s" % (SETTINGS.name, exc))
    if not isinstance(data, dict) or data.get("version") != 1:
        fail("%s 的 version 必须是 1" % SETTINGS.name)
    week1 = data.get("week1")
    if not isinstance(week1, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", week1):
        fail("%s 的 week1 必须是 YYYY-MM-DD" % SETTINGS.name)
    try:
        parsed = date.fromisoformat(week1)
    except ValueError:
        fail("%s 的 week1 不是有效日期" % SETTINGS.name)
    if parsed.isoformat() != week1:
        fail("%s 的 week1 必须零填充且是有效日期" % SETTINGS.name)
    return {"version": 1, "week1": week1}


def read_class_file(path, class_id):
    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        fail("%s 无法读取：%s" % (path.name, exc))
    try:
        if not wb.worksheets:
            fail("%s 没有工作表" % path.name)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            fail("%s 是空工作表" % path.name)

        header_index = None
        header_map = {}
        for row_index, row in enumerate(rows[:20]):
            mapping = {}
            for col, value in enumerate(row):
                name = norm_header(value)
                if name and name not in mapping:
                    mapping[name] = col
            if all(name in mapping for name in REQUIRED):
                header_index = row_index
                header_map = mapping
                break
        if header_index is None:
            fail("%s 找不到表头（需要：%s）" % (path.name, "、".join(REQUIRED)))

        has_date = "日期" in header_map
        has_week = "周次" in header_map
        has_weekday = "星期" in header_map
        if not has_date and not (has_week and has_weekday):
            fail("%s 需要“日期”列，或同时有“周次”和“星期”列" % path.name)

        experiments = []
        for row_number, row in enumerate(rows[header_index + 1:], header_index + 2):
            values = list(row)
            if all(cell_text(v) == "" for v in values):
                continue

            def at(name):
                col = header_map.get(name)
                return values[col] if col is not None and col < len(values) else None

            d = None
            if has_date:
                d = cell_date(at("日期"), wb.epoch)

            # 周次+星期优先；不完整/无法识别时回退到日期。
            week = parse_week(at("周次")) if has_week else None
            weekday = parse_weekday(at("星期")) if has_weekday else None
            if week is not None and weekday is not None:
                d = SETTINGS_WEEK1 + timedelta(days=(week - 1) * 7 + (weekday - 1))

            if d is None:
                fail("%s 第 %d 行：日期无法识别，且周次+星期不完整/无效" % (path.name, row_number))

            period = clean_field(at("节次"), "%s 第 %d 行节次" % (path.name, row_number), MAX_PERIOD_LEN, True)
            name = clean_field(at("实验名称"), "%s 第 %d 行实验名称" % (path.name, row_number), MAX_NAME_LEN, True)
            place = clean_field(at("实验地点"), "%s 第 %d 行实验地点" % (path.name, row_number), MAX_PLACE_LEN, False, "待定")

            experiments.append({
                "date": d.isoformat(),
                "period": period,
                "name": name,
                "place": place,
            })
            if len(experiments) > MAX_EXPERIMENTS:
                fail("%s 的实验数量超过 %d" % (path.name, MAX_EXPERIMENTS))

        return experiments
    finally:
        wb.close()


def scan_class_files():
    if not DATA_DIR.exists():
        fail("找不到 data 目录")
    files = []
    seen = set()
    for path in DATA_DIR.iterdir():
        if not path.is_file():
            continue
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        if path.suffix.lower() != ".xlsx":
            continue
        class_id = path.stem
        if not class_id:
            fail("%s 没有班级名" % path.name)
        if class_id in seen:
            fail("班级名重复：%s" % class_id)
        seen.add(class_id)
        files.append((class_id, path))
    if not files:
        fail("data 目录下没有班级 xlsx")
    files.sort(key=lambda item: natural_key(item[0]))
    return files


def safe_json(data):
    text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    text = text.replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e")
    text = text.replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    return text


def main():
    global SETTINGS_WEEK1

    settings = read_settings()
    SETTINGS_WEEK1 = date.fromisoformat(settings["week1"])

    classes = []
    for class_id, path in scan_class_files():
        experiments = read_class_file(path, class_id)
        classes.append({"id": class_id, "experiments": experiments})

    payload = {"version": 2, "week1": settings["week1"], "classes": classes}

    if not SOURCE.exists():
        fail("找不到 %s" % SOURCE.name)
    html = SOURCE.read_bytes().decode("utf-8-sig").replace("\r\n", "\n")
    matches = list(TEMPLATE_RE.finditer(html))
    if len(matches) != 1:
        fail('源文件中必须恰好有一个 <template id="defaultData">...</template>')

    match = matches[0]
    html = html[: match.start(2)] + safe_json(payload) + html[match.end(2):]
    encoded = html.encode("utf-8")
    SOURCE.write_bytes(encoded)
    DEPLOY.write_bytes(encoded)

    print("built %s: %d classes, %d experiments total" % (
        DEPLOY.name,
        len(classes),
        sum(len(c["experiments"]) for c in classes),
    ))
    for c in classes:
        print("  %s: %d 条" % (c["id"], len(c["experiments"])))

    result = subprocess.run([sys.executable, str(BASE / "verify-csp.py")], cwd=str(BASE))
    if result.returncode != 0:
        fail("verify-csp.py 未通过；请检查源文件/CSP 哈希")
    return 0


if __name__ == "__main__":
    sys.exit(main())
